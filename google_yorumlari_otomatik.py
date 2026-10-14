# -*- coding: utf-8 -*-
"""Windows / Python 3.10+: python google_yorumlari_otomatik.py
Tek dosya. Ilk kullanimda playwright ve openpyxl eksikse pip ile kurulur.
Sifre sadece Chrome'da girilir; OAuth, API anahtari ve BAT gerekmez.
Chrome profili, log, kurtarma verisi ve ciktilar bu dosyanin workspace'inde kalir.
Gorunen Maps arayuzu kullanilir; CAPTCHA/erisim engeli asilmaz.
UI degisirse seciciler guncellenmelidir. %100 icerik garantisi verilmez.
Yerel kontrol: python google_yorumlari_otomatik.py --self-test
Kayitli veriden analiz: python google_yorumlari_otomatik.py --analyze
"""
import importlib.util
import json
import logging
import math
import os
from pathlib import Path
import re
import socket
import subprocess
import sys
import tempfile
import time
import traceback
from collections import Counter
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from urllib.parse import urlsplit, unquote

VERSION = '4.4'
WORKSPACE = Path(__file__).resolve().parent
DEFAULT_OUTPUT = WORKSPACE/'Google_yorum'
CARD = '.jJc9Ad'
CARD_HELPERS = r'''
 const reviewId=c=>c.getAttribute('data-review-id')||c.querySelector('[data-review-id]')?.getAttribute('data-review-id')||'';
 const sections=c=>{
  const body=c.querySelector('.MyEned');let owner=c.querySelector('.CDe7pd');
  if(!owner&&body){
   const outside=[...c.querySelectorAll('.wiI7pd')].filter(n=>!body.contains(n));
   if(outside.length===1)for(let n=outside[0].parentElement;n&&n!==c;n=n.parentElement){
    if(n.querySelectorAll('.wiI7pd').length!==1)break;
    if(n.querySelector('.DZSIDd,[jsaction*="OwnerResponse"]')){owner=n;break;}
   }
  }
  return {body,owner};
 };
 const sectionOf=(c,b)=>{
  const {body,owner}=sections(c);
  if(owner?.contains(b))return '.CDe7pd';
  if(body?.contains(b))return '.MyEned';
  const ids=(b.getAttribute('aria-controls')||'').split(/\s+/).filter(Boolean);
  for(const [selector,n] of [['.CDe7pd',owner],['.MyEned',body]])
   if(n&&[n,...n.querySelectorAll('[id]')].some(x=>x.id&&ids.includes(x.id)))return selector;
  const action=b.getAttribute('jsaction')||'';
  if(owner&&/showOwnerResponseInTranslation/i.test(action))return '.CDe7pd';
  if(body&&/showReviewInTranslation/i.test(action))return '.MyEned';
  return null;
 };
 const isOriginal=b=>b.getAttribute('role')==='switch'&&b.getAttribute('aria-checked')==='true';
 const isExpand=b=>b.getAttribute('aria-expanded')!=='true'&&
  (/expandReview|expandOwnerResponse/.test(b.getAttribute('jsaction')||'')||
   /^(More|See more|Daha fazla(?: göster)?)$/i.test(b.getAttribute('aria-label')||b.textContent.trim()));
 const controls=(c,section,kind)=>[...c.querySelectorAll('button,[role="switch"]')].filter(b=>
  sectionOf(c,b)===section&&(kind==='original'?isOriginal(b):isExpand(b)));
 const sectionText=(c,section)=>{
  const parts=sections(c),n=section==='.MyEned'?parts.body:parts.owner;
  const text=n?.querySelector('.wiI7pd')?.cloneNode(true);if(!text)return null;
  text.querySelectorAll('button,script,style').forEach(x=>x.remove());
  text.querySelectorAll('br').forEach(x=>x.replaceWith('\n'));
  return {text:(text.textContent||'').trim(),language:n.getAttribute('lang')||''};
 };
 const coversText=(full,current)=>!!full?.text&&!!current?.text&&full.language===current.language&&full.text.startsWith(current.text);
 let copyIndex=null;
 const copiesFor=c=>{
  if(!copyIndex){copyIndex=new Map();for(const card of document.querySelectorAll('.jJc9Ad')){
   const id=reviewId(card);if(!copyIndex.has(id))copyIndex.set(id,[]);copyIndex.get(id).push(card);
  }}
  return copyIndex.get(reviewId(c))||[];
 };
 const openEquivalent=(c,section)=>{
  const current=sectionText(c,section);if(!current?.text||controls(c,section,'original').length)return false;
  return copiesFor(c).some(other=>{
   if(other===c||reviewId(other)!==reviewId(c)||controls(other,section,'original').length||controls(other,section,'expand').length)return false;
   const full=sectionText(other,section);
   return coversText(full,current);
  });
 };
 const activeControls=(c,section,kind)=>{
  const pending=controls(c,section,kind);
  return kind==='expand'&&pending.length&&openEquivalent(c,section)?[]:pending;
 };
'''
READ_CARD = r'''e => {
 CARD_HELPERS
 const {body,owner}=sections(e);
 const clean=n=>{if(!n)return '';const c=n.cloneNode(true);
 c.querySelectorAll('button,script,style').forEach(x=>x.remove());
 c.querySelectorAll('br').forEach(x=>x.replaceWith('\n'));return (c.textContent||'').trim();};
 const link=[...e.querySelectorAll('a[href]')].find(a=>/reviewid=|\/review\//.test(a.href));
 const prof=e.querySelector('.al6Kxe');
 const text=clean(body?.querySelector('.wiI7pd')),reply=clean(owner?.querySelector('.wiI7pd'));
 const remaining=[...e.querySelectorAll('.wiI7pd')].filter(n=>!body?.contains(n)&&!owner?.contains(n)).map(clean).filter(Boolean);
 const whole=e.cloneNode(true);whole.querySelectorAll('script,style').forEach(n=>n.remove());
 whole.querySelectorAll('button').forEach(n=>n.querySelector('.d4r55')?n.replaceWith(...n.childNodes):n.remove());
 whole.querySelectorAll('[role="img"][aria-label]').forEach(n=>n.replaceWith(n.getAttribute('aria-label')));
 whole.querySelectorAll('br').forEach(n=>n.replaceWith('\n'));
 whole.querySelectorAll('div,p,li,section').forEach(n=>n.append('\n'));
 const heading=owner?.cloneNode(true);heading?.querySelectorAll('.wiI7pd,.DZSIDd,button,script,style').forEach(n=>n.remove());
 return {id:reviewId(e),
 name:clean(e.querySelector('.d4r55')),star_label:e.querySelector('.kvMYJc')?.getAttribute('aria-label')||'',
 date:clean(e.querySelector('.rsqaWe')),text,language:body?.getAttribute('lang')||body?.querySelector('[lang]')?.getAttribute('lang')||'',
 reply,reply_date:clean(e.querySelector('.DZSIDd')),translated:controls(e,'.MyEned','original').length>0,reply_translated:controls(e,'.CDe7pd','original').length>0,
 translation_unknown:[...e.querySelectorAll('[role="switch"][aria-checked="true"]')].some(b=>!sectionOf(e,b)),
 has_text:!!body,has_reply:!!owner,review_pending:controls(e,'.MyEned','expand').length,reply_pending:controls(e,'.CDe7pd','expand').length,
 all_text:(whole.textContent||'').replace(/\n[ \t]*\n[ \t]*\n/g,'\n\n').trim(),other_text:remaining.join('\n\n'),
 reviewer_info:clean(e.querySelector('.RfnDt')),reply_heading:clean(heading),
 profile:prof?.getAttribute('data-href')||prof?.getAttribute('href')||'',link:link?.href||''};
}'''.replace('CARD_HELPERS',CARD_HELPERS)
READ_META = r'''() => {
 const mains=[...document.querySelectorAll('[role="main"]')].filter(x=>{const r=x.getBoundingClientRect();return r.width>0&&r.height>0;});
 const m=mains.find(x=>x.querySelector('.jJc9Ad'))||mains[0];if(!m)return {};
 const c=m.cloneNode(true);c.querySelectorAll('.jJc9Ad,.jftiEf').forEach(x=>x.remove());
 const labels=[...c.querySelectorAll('[aria-label]')].map(x=>x.getAttribute('aria-label'));
 // Yorum toplami div icinde veya yalniz aria-label'da sunulabilir.
 // Yalniz tam bir sayi+yorum ifadesi kabul edilir; 4.3105 gibi birlesik puan reddedilir.
 const pattern=/^\d+(?:[.,\s]\d{3})*\s+(?:Google\s+)?(?:reviews?|yorum|değerlendirme)(?:\s*\([^)]*\))?$/i;
 // Puan+toplam ust kapsayicisi 4.117,281 gibi sahte bir sayiya donusebilir.
 // Altinda zaten tam bir yorum sayisi etiketi olan kapsayicilari kullanma.
 const texts=[...c.querySelectorAll('span,button,div')].filter(x=>
  ![...x.querySelectorAll('span,button,div')].some(n=>pattern.test((n.textContent||'').trim())))
  .map(x=>(x.textContent||'').trim());
 const counts=[...new Set([...texts,...labels].filter(x=>pattern.test(x)))];
 return {name:m.getAttribute('aria-label')||m.querySelector('h1')?.innerText||document.querySelector('h1')?.innerText||'',counts,labels};
}'''

READ_ACCESS_STATE = r'''() => {
 const visible=e=>{
  if(!e||!e.getClientRects().length)return false;
  for(let p=e;p;p=p.parentElement){const s=getComputedStyle(p);
   if(s.display==='none'||s.visibility==='hidden'||s.visibility==='collapse'||Number(s.opacity)===0)return false;}
  const r=e.getBoundingClientRect();return r.width>0&&r.height>0&&r.bottom>0&&r.right>0&&r.top<innerHeight&&r.left<innerWidth;
 };
 const limited=[],blocked=[];
 const walker=document.createTreeWalker(document.body,NodeFilter.SHOW_TEXT);
 while(walker.nextNode()){
  const n=walker.currentNode,p=n.parentElement;
  if(!p||p.closest('.jJc9Ad,.jftiEf,script,style,noscript,template,[role="tooltip"]')||!visible(p))continue;
  const text=(n.textContent||'').replace(/\s+/g,' ').trim();if(!text)continue;
  const item={text:text.slice(0,800),tag:p.tagName.toLowerCase(),role:p.getAttribute('role')||'',
   context:(p.closest('[role="alert"],[role="status"],[role="dialog"]')||p.parentElement||p).innerText?.slice(0,1000)||''};
  if(/limited view|sınırlı (?:bir )?görünüm/i.test(text)&&!p.closest('a,[role="link"]'))limited.push(item);
  if(/unusual traffic|automated queries|verify you are human|olağandışı trafik|robot olmadığınızı/i.test(text))blocked.push(item);
 }
 return {limited_candidates:limited.slice(0,10),blocked:blocked.length>0,blocked_notices:blocked.slice(0,5),
  captcha:[...document.querySelectorAll('iframe')].some(e=>/recaptcha.*(?:anchor|bframe)|hcaptcha/.test(e.src)&&visible(e)),
  sign_in_visible:[...document.querySelectorAll('a,button,[role="button"]')].some(e=>visible(e)&&/^(Sign in|Oturum aç|Giriş yap)$/i.test((e.innerText||e.getAttribute('aria-label')||'').trim()))};
}'''

def is_limited_notice(text):
 s=re.sub(r'\s+',' ',text).strip().casefold()
 if re.match(r'^(learn more|about |help |sınırlı görünüm hakkında|sınırlı bir görünüm hakkında)',s):return False
 if re.fullmatch(r'limited view(?: of google maps)?|sınırlı (?:bir )?görünüm',s):return True
 if re.match(r'^(limited view|sınırlı (?:bir )?görünüm)\s*:',s):return True
 return bool(re.search(r"(?:you(?:'re|’re| are)|this is|currently|showing|viewing|seeing|using).*limited view",s) or
  re.search(r'sınırlı (?:bir )?görünüm.*(?:kullanıyorsunuz|görüyorsunuz|gösteriliyor|gösterilmektedir|sunuluyor)',s) or
  re.search(r'(?:görüyorsunuz|gösteriliyor|kullanıyorsunuz).*sınırlı (?:bir )?görünüm',s))

def unique_control(detail):
 # Ayni yorumun DOM temsili tekrarlanabilir; fiziksel hedef dugme tek olmali.
 return detail.get('card_count',0)>0 and detail.get('control_count',0)==1

class TaskError(Exception):
 def __init__(self,code,stage,reason,evidence=None):
  self.code,self.stage=code,stage
  self.evidence=evidence
  super().__init__(reason)

def remaining_text(seconds):
 if seconds is None:return 'hesaplanıyor'
 seconds=max(0,int(round(seconds)))
 if seconds<60:return f'{seconds} saniye'
 minutes,seconds=divmod(seconds,60)
 return f'{minutes} dakika'+(f' {seconds} saniye' if seconds else '')

class Progress:
 def __init__(self):
  self.percent=0;self.stage='';self.stage_started=time.monotonic();self.initial_done=0;self.last_line='';self.last_print=0;self.history=[];self.samples=[]
 def update(self,stage,start,end=None,done=None,total=None,note='',remaining=None,force=False):
  now=time.monotonic()
  if stage!=self.stage:self.stage=stage;self.stage_started=now;self.initial_done=done or 0;self.samples=[]
  fraction=min(1,max(0,done/total)) if total and done is not None else 0
  percent=int(start+((end if end is not None else start)-start)*fraction)
  self.percent=max(self.percent,min(100,percent))
  elapsed=now-self.stage_started
  if done is not None:
   if self.samples and done<self.samples[-1][1]:self.samples=[]
   self.samples.append((now,done))
   while len(self.samples)>2 and now-self.samples[1][0]>60:self.samples.pop(0)
  if remaining is None and total and done is not None and self.samples:
   first,initial=self.samples[0];duration=now-first;increase=done-initial
   if duration>=3 and increase>0:remaining=max(0,total-done)*duration/increase
  eta='tamamlanıyor' if total and done is not None and done>=total else remaining_text(remaining)
  count=(f' | {done}/{total}' if total is not None else f' | {done} kayıt alındı; toplam okunamadı') if done is not None else ''
  if stage=='Değerlendirmeler yükleniyor' and total and done is not None:count+=f' | Kayıt ilerlemesi: %{min(100,done/total*100):.1f}'
  line=f'%{self.percent} | {stage}{count}'+(f' | {note}' if note else '')
  if self.percent<100:line+=' | Bu aşamanın tahmini kalanı: '+eta
  if force or line!=self.last_line or now-self.last_print>=5:
   print(line,flush=True);self.last_line=line;self.last_print=now
   self.history.append(dict(stage=stage,percent=self.percent,done=done,total=total,remaining_seconds=round(remaining,1) if remaining is not None else None,at=datetime.now().astimezone().isoformat()))
   logging.info('İLERLEME | %s',line)

def safe_error(e):
 s=re.sub(r'https?://\S+','[URL]',str(e))
 s=re.sub(r'(?i)(token|password|secret|code)=\S+',r'\1=[GIZLENDI]',s)
 return f'{type(e).__name__}: {s[:900]}'

def retry(stage,fn):
 for attempt in range(1,4):
  try:return fn()
  except TaskError:raise
  except Exception as e:
   logging.warning('%s | DENEME %s/3 | %s',stage,attempt,safe_error(e))
   if attempt==3:raise TaskError('UC_DENEME_BASARISIZ',stage,safe_error(e)) from e
   time.sleep(attempt*2)

def bootstrap():
 missing=[p for p in ('playwright','openpyxl') if importlib.util.find_spec(p) is None]
 if missing:
  print('İlk kurulum: eksik paketler kuruluyor:',', '.join(missing))
  retry('PAKET_KURULUMU',lambda:subprocess.run([sys.executable,'-m','pip','install','--index-url','https://pypi.org/simple',*missing],check=True,capture_output=True,text=True))

def validate_url(value):
 p=urlsplit(value)
 host=(p.hostname or '').lower()
 allowed=host in ('maps.app.goo.gl','goo.gl','maps.google.com','maps.google.com.tr','www.google.com','google.com','www.google.com.tr','google.com.tr')
 if p.scheme!='https' or not allowed or p.username or p.password:
  raise TaskError('GECERSIZ_LINK','LINK','https ile baslayan Google Maps isletme baglantisi girin.')
 if host not in ('maps.app.goo.gl','maps.google.com','maps.google.com.tr') and '/maps' not in p.path:
  raise TaskError('GECERSIZ_LINK','LINK','Baglanti Google Maps baglantisi olmali.')
 return value

def map_place_key(value):
 path=unquote(urlsplit(value).path)
 match=re.search(r'!1s([^!/]+)',path)
 if match:return match[1]
 match=re.search(r'/maps/place/([^/]+)',path)
 return 'name:'+match[1].replace('+',' ').casefold() if match else None

def metadata(raw):
 totals=set()
 for s in raw.get('counts',[]):
  m=re.fullmatch(r'(\d+|\d{1,3}(?:[.,\s]\d{3})+)\s+(?:Google\s+)?(?:reviews?|yorum|değerlendirme)(?:\s*\([^)]*\))?',s.strip(),re.I)
  if m:totals.add(int(re.sub(r'\D','',m[1])))
 hist={};rating=None
 for label in raw.get('labels',[]):
  m=re.search(r'([1-5])\s*(?:yıldızlı|stars?)[,\s]+([\d\s.,]+)\s*(?:yorum|reviews?)',label,re.I)
  if m:hist[int(m[1])]=int(re.sub(r'\D','',m[2]))
  m=re.fullmatch(r'\s*([1-5](?:[.,]\d+)?)\s*(?:yıldızlı|stars?)\s*',label,re.I)
  if m:rating=float(m[1].replace(',','.'))
 histogram_total=sum(hist.values()) if len(hist)==5 else None
 if len(totals)==1:total=next(iter(totals));source='baslik'
 elif histogram_total is not None and histogram_total in totals:total=histogram_total;source='baslik ve yildiz dagilimi'
 elif not totals and histogram_total is not None:total=histogram_total;source='yildiz dagilimi'
 else:total=None;source='belirsiz'
 return dict(name=raw.get('name',''),total=total,rating=rating,histogram=hist,
  total_source=source,total_candidates=sorted(totals),histogram_total=histogram_total)

def merge_raw_variants(items):
 # Bir anlik DOM taramasindaki ayni ID temsilleri; yorum ve yanit bagimsiz secilir.
 base=dict(max(items,key=lambda x:sum(bool(x.get(k)) for k in ('name','star_label','date','profile','text','reply'))))
 conflicts=[]
 for section,fields,translated,pending in (
  ('has_text',('text','language','review_pending','translated'),'translated','review_pending'),
  ('has_reply',('reply','reply_date','reply_heading','reply_pending','reply_translated'),'reply_translated','reply_pending')):
  candidates=[x for x in items if x.get(section)]
  if not candidates:continue
  originals=[x for x in candidates if not x.get(translated)]
  chosen=max(originals or candidates,key=lambda x:(not x.get(pending),len(x.get(fields[0],''))))
  base[section]=True
  for key in fields:base[key]=chosen.get(key,False if key==translated else 0 if key==pending else '')
  texts={x.get(fields[0],'') for x in originals if x.get(fields[0])}
  if any(not (a.startswith(b) or b.startswith(a)) for a in texts for b in texts):conflicts.append(fields[0])
  translated_texts=[x.get(fields[0],'') for x in candidates if x.get(translated)]
  if translated_texts:base['saved_'+translated]=max(translated_texts,key=len)
 for key in ('name','star_label','date','profile','link','reviewer_info'):
  if not base.get(key):base[key]=next((x[key] for x in items if x.get(key)),'')
 base['all_text']='\n\n'.join(dict.fromkeys(x['all_text'] for x in items if x.get('all_text')))
 base['other_text']='\n\n'.join(dict.fromkeys(x['other_text'] for x in items if x.get('other_text')))
 base['variant_count']=len(items);base['variant_conflicts']=conflicts
 base['translation_unknown']=any(x.get('translation_unknown') for x in items)
 return base

def normalize(raw):
 m=re.search(r'(?<!\d)([1-5])(?:[.,]0)?\s*(?:stars?|yıldız)',raw.get('star_label',''),re.I)
 star=int(m[1]) if m else None
 issues=[]
 for key,label in [('id','yorum kimligi'),('name','kullanici adi'),('date','yorum tarihi')]:
  if not raw.get(key):issues.append(label+' okunamadi')
 if star is None:issues.append('yildiz okunamadi')
 for key,label in [('review_pending','yorum'),('reply_pending','yanit')]:
  if raw.get(key):issues.append(label+' genisletilemedi')
 if raw.get('translated'):issues.append('ozgun yorum gosterilemedi')
 if raw.get('reply_translated'):issues.append('ozgun isletme yaniti gosterilemedi')
 if raw.get('translation_unknown'):issues.append('ceviri dugmesinin ait oldugu metin belirlenemedi')
 if raw.get('other_text'):issues.append('ek metin var; yorum/yanit alanina eslenemedi (tum metinlerde korundu)')
 if raw.get('variant_conflicts'):issues.append('ayni yorumun DOM kopyalarinda metin farki var: '+', '.join(raw['variant_conflicts']))
 if raw.get('retained_text'):issues.append('yorum onceki okumadan korundu; son gorunumde metin okunamadi')
 if raw.get('retained_reply'):issues.append('yanit onceki okumadan korundu; son gorunumde metin okunamadi')
 text=raw.get('text','');reply=raw.get('reply','')
 if raw.get('has_text') and not text:issues.append('yorum alani var ama bos')
 if raw.get('has_reply') and not reply:issues.append('yanit alani var ama bos')
 if raw.get('has_reply') and reply and not raw.get('reply_date'):issues.append('yanit tarihi okunamadi')
 # Noktalama yazara aittir. Eksiklik yalniz DOM kontrol kanitiyla belirlenir.
 row={k:raw.get(k,'') for k in ('id','name','date','language','reply_date','profile','link','all_text','other_text','reviewer_info','reply_heading')}
 row.update(stars=star,text='' if raw.get('translated') else text,
  translation=text if raw.get('translated') else raw.get('saved_translated',''),reply='' if raw.get('reply_translated') else reply,
  reply_translation=reply if raw.get('reply_translated') else raw.get('saved_reply_translated',''),
  comment_type='METINLI YORUM' if text else ('DOGRULANAMADI' if raw.get('has_text') or raw.get('other_text') else 'YALNIZCA YILDIZ; YORUM YAZILMAMIS'),
  quality='ALAN KONTROLU TAM' if not issues else 'EKSIK / DOGRULANAMADI',issues='; '.join(issues),
  checks={k:raw.get(k,False if 'translat' in k else 0) for k in ('review_pending','reply_pending','translated','reply_translated','translation_unknown')})
 row['review_state']=('METIN ACIK / OKUNDU' if text and not raw.get('review_pending') and not raw.get('translated') else 'METIN TAM / OZGUN DOGRULANAMADI') if raw.get('has_text') else ('EK METIN VAR; YORUM DURUMU INCELENMELI' if raw.get('other_text') else 'BU GORUNUMDE YALNIZCA YILDIZ')
 row['reply_state']=('YANIT ACIK / OKUNDU' if reply and not raw.get('reply_pending') and not raw.get('reply_translated') else 'YANIT VAR; TAM / OZGUN DOGRULANAMADI') if raw.get('has_reply') else ('EK METIN VAR; YANIT DURUMU INCELENMELI' if raw.get('other_text') else 'BU GORUNUMDE ISLETME YANITI YOK')
 if raw.get('retained_text'):row['review_state']='ONCEKI OKUMADAN KORUNDU; SON GORUNUM DOGRULANAMADI'
 if raw.get('retained_reply'):row['reply_state']='ONCEKI OKUMADAN KORUNDU; SON GORUNUM DOGRULANAMADI'
 return row

def age(label):
 s=str(label).lower().strip()
 m=re.search(r'(?:(\d+)|\b(a|an|bir))\s+(minutes?|hours?|days?|weeks?|months?|years?|dakika|saat|gün|hafta|ay|yıl)',s)
 if m:
  n=int(m[1]) if m[1] else 1
  unit=m[3].rstrip('s')
  return n*{'minute':1/1440,'dakika':1/1440,'hour':1/24,'saat':1/24,'day':1,'gün':1,'week':7,'hafta':7,'month':30.436875,'ay':30.436875,'year':365.2425,'yıl':365.2425}[unit]
 if s in ('just now','az önce','şimdi','today','bugün'):return 0
 if s in ('yesterday','dün'):return 1
 return float('inf')

def checked_page(page):
 if page.is_closed():raise TaskError('TARAYICI_KAPANDI','ERISIM','Chrome sekmesi kapatildi.')
 if '/sorry/' in page.url:raise TaskError('ERISIM_ENGELI','ERISIM','Google robot kontrolu gosteriyor; islem durdu.')
 p=urlsplit(page.url)
 if p.hostname=='accounts.google.com':raise TaskError('GIRIS_TAMAMLANMADI','GIRIS','Google girisi tamamlanmadi; Chrome ekranindaki bildirimi kontrol edin.')
 state=page.evaluate(READ_ACCESS_STATE)
 state['limited_notices']=[n for n in state['limited_candidates'] if is_limited_notice(n['text'])]
 state['limited']=bool(state['limited_notices'])
 if state['blocked'] or state['captcha']:raise TaskError('ERISIM_ENGELI','ERISIM','CAPTCHA/erisim engeli goruldu. Engel asilmaz.',evidence=state)
 return state

def chrome(pw):
 if os.name!='nt':raise TaskError('WINDOWS_GEREKLI','CHROME','Bu dosya Windows ve kurulu Google Chrome icin hazirlandi.')
 paths=[Path(os.environ.get(k,'C:/'))/'Google/Chrome/Application/chrome.exe' for k in ('PROGRAMFILES','PROGRAMFILES(X86)','LOCALAPPDATA')]
 binary=next((p for p in paths if p.is_file()),None)
 if not binary:raise TaskError('CHROME_YOK','CHROME','Google Chrome kurulu degil veya standart konumda bulunamadi.')
 profile=WORKSPACE/'Chrome_profili'
 profile.mkdir(parents=True,exist_ok=True)
 logging.info('Chrome profili: %s',profile)
 with socket.socket() as sock:
  sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
 subprocess.Popen([str(binary),f'--user-data-dir={profile}',f'--remote-debugging-port={port}','--remote-debugging-address=127.0.0.1','--no-first-run','about:blank'])
 endpoint=f'http://127.0.0.1:{port}'
 time.sleep(2)
 def attach():
  try:return pw.chromium.connect_over_cdp(endpoint,timeout=12000)
  except Exception as e:raise RuntimeError('Chrome baglantisi kurulamadi. Bu uygulamanin onceki Chrome penceresi aciksa kapatip yeniden calistirin. '+safe_error(e)) from e
 return retry('CHROME_BAGLANTISI',attach)

class Collector:
 def __init__(self,url,run_dir=None):
  self.url=url;self.rows={};self.raw_rows={};self.meta={};self.final_meta={};self.error=None;self.round=0
  self.access_state={};self.access_screenshot=''
  self.progress=Progress()
  self.performance=dict(snapshot_calls=0,dom_read_seconds=0.0,merge_seconds=0.0,checkpoint_calls=0,checkpoint_seconds=0.0)
  self.diagnostics=[]
  self.started=datetime.now().astimezone().isoformat()
  self.deadline=time.monotonic()+7200
  self.checkpoint_dir=Path(run_dir) if run_dir else WORKSPACE/'Calisma_kayitlari'/datetime.now().strftime('%Y%m%d_%H%M%S_%f')
  self.checkpoint_dir.mkdir(parents=True,exist_ok=True)
  self.checkpoint_file=self.checkpoint_dir/'yorum_kurtarma.json'
 def checkpoint(self):
  measured=time.perf_counter()
  data=dict(version=VERSION,url=self.url,rows=list(self.rows.values()),raw_rows=self.raw_rows,meta=self.meta,final_meta=self.final_meta,error=self.error,diagnostics=self.diagnostics,access_state=self.access_state,access_screenshot=self.access_screenshot,progress=self.progress.history,performance=self.performance,started=self.started,round=self.round)
  tmp=self.checkpoint_file.with_suffix('.tmp')
  tmp.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8');tmp.replace(self.checkpoint_file)
  self.performance['checkpoint_calls']+=1
  self.performance['checkpoint_seconds']+=time.perf_counter()-measured
 def check_access(self,page,allow_limited=False):
  try:state=checked_page(page)
  except TaskError as e:
   if e.evidence:self.access_state=e.evidence;self.capture_access(page)
   raise
  self.access_state=state
  if state['limited'] and not allow_limited:
   # Yeni gorunen kayitlari ve bildirimi durmadan once koru.
   self.snapshot(page);self.capture_access(page);self.checkpoint()
   notice=' | '.join(n['text'] for n in state['limited_notices'])
   logging.error('GORUNUR SINIRLI GORUNUM | %s | GIRIS DUGMESI=%s',notice,state['sign_in_visible'])
   raise TaskError('SINIRLI_GORUNUM','TOPLAMA','Gorunur sinirli gorunum bildirimi: '+notice+
    '. Chrome ekraninda Google girisini ve bildirimi kontrol edip yeniden calistirin.',evidence=state)
  return state
 def capture_access(self,page):
  try:
   path=self.checkpoint_dir/'erisim_bildirimi.png'
   page.screenshot(path=str(path));self.access_screenshot=str(path)
  except Exception as e:logging.warning('ERISIM EKRANI KAYDEDILEMEDI | %s',safe_error(e))
 def locator(self,page,rid):
  # Tek bir kart lokatoru; alt ogelerde tekrar eden ID'ler karti cogaltmaz.
  attr='[data-review-id='+json.dumps(rid)+']'
  return page.locator(CARD+attr+', '+CARD+':has('+attr+')').first
 def store(self,raw):
  if not raw.get('id'):raise TaskError('KIMLIK_YOK','YORUM_OKU','Google yorum kimligi yok; ayni adli kullanicilarin karismamasi icin duruldu.')
  raw=dict(raw);previous=self.raw_rows.get(raw['id'],{})
  for present,fields,flag in (
   ('has_text',('text','language','review_pending','translated'),'retained_text'),
   ('has_reply',('reply','reply_date','reply_heading','reply_pending','reply_translated'),'retained_reply')):
   if not raw.get(fields[0]) and previous.get(fields[0]):
    raw[present]=True;raw[flag]=True
    for key in fields:raw[key]=previous.get(key,'')
  if raw.get('retained_text') or raw.get('retained_reply'):
   raw['all_text']='\n\n'.join(dict.fromkeys(x for x in (raw.get('all_text',''),previous.get('all_text','')) if x))
  self.raw_rows[raw['id']]=raw
  row=normalize(raw);old=self.rows.get(row['id'])
  if old:
   for key in ('translation','reply_translation'):
    if not row[key]:row[key]=old.get(key,'')
  # Son DOM kontrolu esas: eski TAM etiketi yeni kapali dugmeyi gizleyemez.
  self.rows[row['id']]=row
 def read_one(self,page,rid):
  self.store(self.locator(page,rid).evaluate(READ_CARD));self.checkpoint()
  self.expand_remaining(page,rid=rid)
  # Ozgune donus yeni bir devam dugmesi gosterebilir.
  self.expand_remaining(page,rid=rid)
  self.store(self.locator(page,rid).evaluate(READ_CARD));self.checkpoint()
  row=self.rows[rid]
  if row['quality']!='ALAN KONTROLU TAM':
   logging.warning('ALAN KONTROLU | %s | %s',row['name'],row['issues'])
   raise TaskError('YORUM_ALANI_EKSIK','YORUM_OKU',f'{row["name"]}: '+row['issues'])
 def visible(self,page):
  ids=page.locator(CARD).evaluate_all("es=>es.map(e=>e.getAttribute('data-review-id')||e.querySelector('[data-review-id]')?.getAttribute('data-review-id')||'')")
  for rid in dict.fromkeys(ids):
   if not rid:raise TaskError('KIMLIK_YOK','YORUM_OKU','Kimligi olmayan yorum karti bulundu.')
   if rid not in self.rows:
    retry('YORUM_OKU',lambda rid=rid:self.read_one(page,rid))
 def prepare_reviews(self,page):
  self.check_access(page)
  mains=[m for m in page.get_by_role('main').all() if m.is_visible()]
  if len(mains)!=1:raise TaskError('YORUM_PANELI_BELIRSIZ','YORUMLAR_AC','Tek bir gorunur isletme paneli bulunamadi.')
  panel=mains[0]
  tabs=[t for t in panel.get_by_role('tab',name=re.compile('Yorumlar|Reviews',re.I)).all() if t.is_visible()]
  if len(tabs)>1:raise TaskError('YORUM_SEKMESI_BELIRSIZ','YORUMLAR_AC','Birden fazla yorum sekmesi bulundu.')
  if tabs:
   if tabs[0].get_attribute('aria-selected')!='true':retry('YORUMLAR_AC',lambda:tabs[0].click(timeout=15000))
   retry('YORUM_SEKMESI_SECIMI',lambda:panel.get_by_role('tab',name=re.compile('Yorumlar|Reviews',re.I),selected=True).wait_for(state='visible',timeout=10000))
   if tabs[0].get_attribute('aria-selected')!='true':raise TaskError('YORUM_SEKMESI_ACILMADI','YORUMLAR_AC','Yorum sekmesi tiklandi ancak secili degil.')
  else:
   buttons=[b for b in panel.get_by_role('button',name=re.compile(r'^\s*\d[\d.,\s]*\s+(?:Google\s+)?(?:yorum|reviews?|değerlendirme)\s*$',re.I)).all() if b.is_visible()]
   if len(buttons)==1:retry('YORUMLAR_AC',lambda:buttons[0].click(timeout=15000))
   elif not page.locator(CARD).count():raise TaskError('YORUM_GIRISI_BULUNAMADI','YORUMLAR_AC','Yorum sekmesi veya degerlendirme dugmesi bulunamadi.')
  self.meta=metadata(retry('ILK_TOPLAM_OKUMA',lambda:page.evaluate(READ_META)));self.checkpoint()
  if self.meta.get('total')==0:return
  self.progress.update('İlk yorumlar yükleniyor',10)
  # Honda gibi sayfalarda filtrelerin altindaki kartlar panel kaydirilinca yuklenir.
  # Kaydirma alani ilk karta bagli olmadan bulunur. Sinirli ve yeniden kontrol edilen kurtarma.
  panel_script=r'''e=>{
   const visible=n=>{const r=n.getBoundingClientRect();return r.width>0&&r.height>0;};
   const nodes=[e,...e.querySelectorAll('div')].filter(n=>visible(n)&&n.clientHeight>80&&n.scrollHeight>n.clientHeight+5&&/auto|scroll/.test(getComputedStyle(n).overflowY));
   const selected=nodes.sort((a,b)=>b.clientHeight-a.clientHeight)[0];
   if(!selected)return {found:false};
   const before=selected.scrollTop;selected.scrollTop+=Math.max(250,selected.clientHeight*.85);
   return {found:true,before,after:selected.scrollTop,height:selected.scrollHeight};
  }'''
  for attempt in range(1,7):
   self.check_access(page)
   visible=[card for card in page.locator(CARD).all() if card.is_visible()]
   if visible:
    self.meta=metadata(page.evaluate(READ_META));self.checkpoint();return
   movement=retry('ILK_PANEL_KAYDIR',lambda:panel.evaluate(panel_script))
   logging.info('ILK_YORUM_KURTARMA | DENEME=%s | HAREKET=%s',attempt,movement)
   page.wait_for_timeout(1500 if movement.get('found') else 2500)
  if any(card.is_visible() for card in page.locator(CARD).all()):
   self.meta=metadata(page.evaluate(READ_META));self.checkpoint();return
  evidence=dict(meta=self.meta,url=page.url,attempts=6)
  path=self.checkpoint_dir/'ilk_yorum_yuklenmedi.png'
  try:page.screenshot(path=str(path));evidence['screenshot']=str(path)
  except Exception as e:logging.warning('ILK_YORUM_EKRANI | %s',safe_error(e))
  self.diagnostics.append(dict(stage='ILK_YORUM',**evidence))
  raise TaskError('ILK_YORUM_YUKLENMEDI','ILK_YORUM','Yorum paneli acildi ancak kontrollu kaydirmaya ragmen kartlar yuklenmedi.',evidence=evidence)
 def collect(self,page):
  self.check_access(page,allow_limited=True)
  self.prepare_reviews(page)
  if not self.meta.get('name'):raise TaskError('YER_ADI_OKUNAMADI','KONUM','Isletme adi okunamadi; yanlis sayfa olabilir.')
  logging.info('ISLETME=%s | TOPLAM=%s | ORTALAMA=%s',self.meta['name'],self.meta['total'],self.meta['rating'])
  logging.info('TOPLAM KAYNAGI | %s | ADAYLAR=%s | HISTOGRAM=%s',self.meta.get('total_source'),self.meta.get('total_candidates'),self.meta.get('histogram_total'))
  if self.meta.get('total')==0:
   self.final_meta=dict(self.meta);self.checkpoint();return
  # Maps menulerinin rol/dil farki toplamayi durdurmamali.
  # Tum kayitlar mevcut siralamada toplanir; Excel tarihi yeniden eskiye siralar.
  retry('ACIK_MENU_KAPAT',lambda:page.keyboard.press('Escape'))
  logging.info('SIRALAMA | Maps mevcut siralamasi korunuyor; Excel yeniden eskiye siralanacak.')
  if self.meta.get('total') is None:
   logging.warning('TOPLAM OKUNAMADI | Toplama devam edecek; sayi bilinmiyorsa tam sonuc kabul edilmeyecek.')
  self.progress.update('Değerlendirmeler yükleniyor',10,60,done=0,total=self.meta.get('total'))
  idle=0
  short_stalls=0
  for turn in range(1,10001):
   self.round=turn
   if time.monotonic()>self.deadline:raise TaskError('SURE_SINIRI','TOPLAMA','Iki saatlik guvenlik sinirina ulasildi.')
   self.check_access(page);before=len(self.rows)
   # Asama 1: hizli yukleme. Metinleri genisletme, sadece ekrandaki kayitlari koru.
   self.snapshot(page)
   def scroll():
    movement=page.locator(CARD).first.evaluate(r'''e=>{for(let n=e.parentElement;n;n=n.parentElement){
     if(n.scrollHeight>n.clientHeight+5&&/auto|scroll/.test(getComputedStyle(n).overflowY)){
      const b=n.scrollTop;n.scrollTop=n.scrollHeight;
      return {before:b,after:n.scrollTop,height:n.scrollHeight,bottom:n.scrollTop+n.clientHeight>=n.scrollHeight-5};}}
     return null;}''')
    if movement is None:raise RuntimeError('Yorum panelinin kaydirilabilir alani bulunamadi.')
    return movement
   movement=retry('YORUM_PANELI_KAYDIR',scroll)
   # Yeni kimlik DOM'a geldiginde hemen devam et; ag gecikmesinde 4 saniyeye kadar bekle.
   # Ayni yorumun kopyasi veya metninin degismesi yeni kayit sayilmaz.
   from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
   waiting=time.perf_counter();arrived=False
   try:
    page.wait_for_function(r'''known=>{
     const seen=new Set(known);
     return [...document.querySelectorAll('.jJc9Ad')].some(c=>{
      const id=c.getAttribute('data-review-id')||c.querySelector('[data-review-id]')?.getAttribute('data-review-id');
      return id&&!seen.has(id);
     });
    }''',arg=list(self.rows),timeout=4000,polling=200)
    arrived=True
   except PlaywrightTimeoutError:pass
   logging.info('YENI_KAYIT_BEKLEME | YENI_KIMLIK=%s | SANIYE=%.3f',arrived,time.perf_counter()-waiting)
   self.check_access(page);self.snapshot(page)
   # Yorum paneli yuklenirken baslik sonradan gelebilir.
   if self.meta.get('total') is None:
    latest=metadata(page.evaluate(READ_META))
    if latest.get('total') is not None:
     self.meta.update(latest)
     logging.info('TOPLAM SONRADAN OKUNDU | %s',latest['total'])
   logging.info('YUKLEME TUR=%s | BULUNAN=%s | MAPS=%s | Kaydirma=%s',turn,len(self.rows),self.meta.get('total'),movement)
   self.progress.update('Değerlendirmeler yükleniyor',10,60,done=len(self.rows),total=self.meta.get('total'))
   # Hic yeni kayit gelmemesi ancak liste dibinde bitis adayi sayilir.
   idle=idle+1 if len(self.rows)==before and movement['bottom'] else 0
   if len(self.rows)>before:short_stalls=0
   if idle:
    logging.warning('YENI KAYIT GELMEDI | Liste dibi kontrolu %s/3',idle)
    self.progress.update('Değerlendirmeler yükleniyor',10,60,done=len(self.rows),total=self.meta.get('total'),note='Yeni kayıt bekleniyor; bitiş henüz doğrulanmadı')
   if idle>=3:
    # Bitis adayi sonrasi bir kez daha bekle: gec gelen kayitlari atlama.
    page.wait_for_timeout(4000);self.check_access(page)
    last=len(self.rows);self.snapshot(page)
    if len(self.rows)>last:idle=0;short_stalls=0;continue
    # Bilinen toplamdan az kayitta kisa bir ag duraklamasini bitis sayma.
    target=self.meta.get('total')
    if target is not None and len(self.rows)<target:
     short_stalls+=1
     if short_stalls<3:
      logging.warning('TOPLAM ALTINDA DURAKLAMA | %s/%s | KURTARMA=%s/3',len(self.rows),target,short_stalls)
      retry('YUKLEME_YENIDEN_TETIKLE',lambda:page.locator(CARD).first.evaluate(r'''e=>{
       for(let n=e.parentElement;n;n=n.parentElement)if(n.scrollHeight>n.clientHeight+5&&/auto|scroll/.test(getComputedStyle(n).overflowY)){
        n.scrollTop=Math.max(0,n.scrollTop-Math.max(200,n.clientHeight*.6));return true;
       }return false;
      }'''))
      page.wait_for_timeout(5000);idle=0;continue
     raise TaskError('YORUM_YUKLEME_ERKEN_DURDU','TOPLAMA',f'Gorunen toplam {target}, alinan {len(self.rows)}. Yeni yorumlar sinirli yeniden denemelerde yuklenmedi.')
    logging.info('YUKLEME BITTI | %s benzersiz kayit.',len(self.rows))
    self.progress.update('Yorumlar ve işletme cevapları açılıyor',65)
    self.expand_bulk(page)
    self.snapshot(page)
    self.final_meta=metadata(page.evaluate(READ_META));self.checkpoint();return
  raise TaskError('TUR_SINIRI','TOPLAMA','10000 kaydirma sinirina ulasildi.')
 def snapshot(self,page):
  begin=time.perf_counter()
  script='es=>es.map('+READ_CARD+')'
  raw=retry('YUKLENEN_KARTLARI_KORU',lambda:page.locator(CARD).evaluate_all(script))
  read_end=time.perf_counter()
  grouped={}
  for item in raw:grouped.setdefault(item.get('id',''),[]).append(item)
  for items in grouped.values():self.store(merge_raw_variants(items))
  merge_end=time.perf_counter()
  self.performance['snapshot_calls']+=1
  self.performance['dom_read_seconds']+=read_end-begin
  self.performance['merge_seconds']+=merge_end-read_end
  self.checkpoint()
  logging.info('OKUMA_SURELERI | KAYIT=%s | DOM=%.3f | BIRLESTIRME=%.3f | JSON=%.3f',len(self.rows),read_end-begin,merge_end-read_end,time.perf_counter()-merge_end)
 def expand_loaded(self,page):
  # Maps gorunumden cikmis kartlari DOM'dan silerse, ham kayit korunur ancak tam sayilmaz.
  loaded=set(page.locator(CARD).evaluate_all("es=>es.map(e=>e.getAttribute('data-review-id')||e.querySelector('[data-review-id]')?.getAttribute('data-review-id')||'')"))
  absent=set(self.rows)-loaded
  for i,rid in enumerate(list(self.rows),1):
   if rid in absent:continue
   logging.info('METIN OKUMA | %s/%s | %s',i,len(self.rows),self.rows[rid]['name'])
   self.check_access(page)
   retry('YORUM_OKU',lambda rid=rid:self.read_one(page,rid))
  if absent:
   raise TaskError('KARTLAR_DOMDAN_KALDIRILDI','METIN_OKU',f'{len(absent)} kayit kaydirma sirasinda sayfadan kaldirildi. Ham gorunumleri korundu; tam metinleri dogrulanamadi.')
 def expand_bulk(self,page):
  # Degisen dugme dizinlerini beklemek yerine her partide canli DOM yeniden taranir.
  # Gorunen arayuzun Daha fazla ve Ozgun metin dugmeleri kullanilir; gizli API yoktur.
  script=r'''limit=>{
   CARD_HELPERS
   const cards=[...document.querySelectorAll('.jJc9Ad')];
   const original=[...new Set(cards.flatMap(c=>['.MyEned','.CDe7pd'].flatMap(s=>activeControls(c,s,'original'))))];
   const expand=[...new Set(cards.flatMap(c=>['.MyEned','.CDe7pd'].flatMap(s=>activeControls(c,s,'expand'))))];
   const pending=original.length?original:expand;
   let clicked=0;for(const b of pending.slice(0,limit))if(b.isConnected){b.click();clicked++;}
   return {original:original.length,expand:expand.length,clicked};
  }'''.replace('CARD_HELPERS',CARD_HELPERS)
  self.bulk_script=script
  stalled=0;previous=None;initial_pending=0
  for batch in range(1,2001):
   if time.monotonic()>self.deadline:raise TaskError('SURE_SINIRI','METIN_GENISLET','Iki saatlik guvenlik sinirina ulasildi.')
   self.check_access(page)
   self.snapshot(page) # Ceviriler ozgune gecmeden once de korunur.
   state=retry('TOPLU_METIN_GENISLET',lambda:page.evaluate(script,12))
   pending=state['original']+state['expand']
   initial_pending=max(initial_pending,pending)
   self.progress.update('Yorumlar ve işletme cevapları açılıyor',65,90,done=max(0,initial_pending-pending),total=initial_pending or 1)
   logging.info('METIN GENISLETME | PARTI=%s | OZGUNE DONUS=%s | KAPALI DAHA FAZLA=%s | TIKLANAN=%s',batch,state['original'],state['expand'],state['clicked'])
   if pending==0:
    page.wait_for_timeout(600);self.snapshot(page)
    final=page.evaluate(script,0)
    if final['original']+final['expand']==0:break
    continue
   signature=(state['original'],state['expand'])
   stalled=stalled+1 if signature==previous else 0
   previous=signature
   if stalled>=3:
    logging.warning('TOPLU ACMA DURAKLADI | %s kontrol kaldi; normal gorunur tiklama denenecek.',pending)
    self.expand_remaining(page)
    stalled=0;previous=None
   page.wait_for_timeout(400 if not stalled else 2000)
  else:raise TaskError('GENISLETME_SINIRI','METIN_GENISLET','2000 parti sinirina ulasildi.')
  loaded=set(page.locator(CARD).evaluate_all("es=>es.map(e=>e.getAttribute('data-review-id')||e.querySelector('[data-review-id]')?.getAttribute('data-review-id')||'')"))
  absent=set(self.rows)-loaded
  if absent:raise TaskError('KARTLAR_DOMDAN_KALDIRILDI','METIN_GENISLET',f'{len(absent)} yorum DOMdan kaldirildi; kayitlar korunuyor ama tamlik dogrulanamadi.')
  problems=[]
  for row in self.rows.values():
   serious=[x for x in row['issues'].split('; ') if x]
   if serious:problems.append(row['name']+': '+ '; '.join(serious))
  if problems:raise TaskError('YORUM_ALANI_EKSIK','METIN_GENISLET',' | '.join(problems[:5]))
  logging.info('GENISLETME BITTI | Kapali devam / etkin ceviri kontrolu kalmadi. Noktalama birebir korunuyor.')
  self.progress.update('Yorum ve cevaplar kontrol edildi',90)
 def pending_controls(self,page):
  return page.locator(CARD).evaluate_all(r'''es=>es.flatMap(c=>{
   CARD_HELPERS
   const id=reviewId(c);
   const name=c.querySelector('.d4r55')?.textContent||'';const out=[];
   for(const section of ['.MyEned','.CDe7pd']){
    for(const kind of ['original','expand'])if(activeControls(c,section,kind).length)out.push({id,name,section,kind});
   }return out;
  })'''.replace('CARD_HELPERS',CARD_HELPERS))
 def prepare_control(self,page,target,token):
  return page.evaluate(r'''arg=>{
   CARD_HELPERS
   document.querySelectorAll('[data-google-yorum-target]').forEach(b=>b.removeAttribute('data-google-yorum-target'));
   const cards=[...document.querySelectorAll('.jJc9Ad')].filter(c=>reviewId(c)===arg.target.id);
   const buttons=[...new Set(cards.flatMap(c=>activeControls(c,arg.target.section,arg.target.kind)))];
   const b=buttons[0];
   const detail={card_count:cards.length,control_count:buttons.length};
   if(!b)return detail;
   b.setAttribute('data-google-yorum-target',arg.token);
   const rect=b.getBoundingClientRect(),style=getComputedStyle(b);
   return {...detail,text:b.textContent.trim(),aria_label:b.getAttribute('aria-label'),
    aria_expanded:b.getAttribute('aria-expanded'),aria_checked:b.getAttribute('aria-checked'),
    action:b.getAttribute('jsaction'),visible:rect.width>0&&rect.height>0&&style.display!=='none'&&style.visibility!=='hidden',
    in_viewport:rect.bottom>0&&rect.right>0&&rect.top<innerHeight&&rect.left<innerWidth,
    disabled:!!b.disabled,rect:{x:rect.x,y:rect.y,width:rect.width,height:rect.height}};
  }'''.replace('CARD_HELPERS',CARD_HELPERS),dict(target=target,token=token))
 def control_evidence(self,page,target):
  script=r'''target=>{
   CARD_HELPERS
   const read=READ_FUNCTION;
   return [...document.querySelectorAll('.jJc9Ad')].filter(c=>reviewId(c)===target.id).map(c=>{
    const raw=read(c);return {raw,markup:c.outerHTML.slice(0,40000)};
   });
  }'''.replace('CARD_HELPERS',CARD_HELPERS).replace('READ_FUNCTION',READ_CARD)
  return page.evaluate(script,target)
 def expand_remaining(self,page,rid=None):
  targets=[x for x in self.pending_controls(page) if rid is None or x['id']==rid]
  # Ceviriden donmeden devam metnini acmak eski hedefi gecersiz kilabilir.
  targets.sort(key=lambda x:x['kind']!='original')
  self.progress.update('Kalan metinler kontrol ediliyor',88,91,done=0,total=len(targets) or 1)
  failed=[]
  # Sadece takilan kartlar islenir; tum yorumlar tekrar tek tek okunmaz.
  for target_index,target in enumerate(targets,1):
   self.progress.update('Kalan metinler kontrol ediliyor',88,91,done=target_index-1,total=len(targets),note=target['name'])
   for attempt in range(1,4):
    if time.monotonic()>self.deadline:raise TaskError('SURE_SINIRI','METIN_GENISLET','Iki saatlik guvenlik sinirina ulasildi.')
    self.check_access(page)
    current=[x for x in self.pending_controls(page) if x['id']==target['id'] and x['section']==target['section'] and x['kind']==target['kind']]
    if not current:break
    detail={};before=[]
    try:
     before=self.control_evidence(page,target)
     token=f'control-{time.monotonic_ns()}'
     detail=self.prepare_control(page,target,token)
     if not unique_control(detail):raise RuntimeError('Hedef dugme bulunamadi veya tek bir dugme olarak belirlenemedi: '+json.dumps(detail,ensure_ascii=False))
     # Okuyucuyla ayni kuralla bulunan canli dugme isaretlenir. Her denemede yeniden bulunur.
     button=page.locator('[data-google-yorum-target='+json.dumps(token)+']')
     button.scroll_into_view_if_needed(timeout=6000)
     button.click(timeout=4000)
     page.wait_for_function(r'''target=>{
      CARD_HELPERS
      const cards=[...document.querySelectorAll('.jJc9Ad')].filter(c=>reviewId(c)===target.id);
      return cards.length>0&&cards.every(c=>activeControls(c,target.section,target.kind).length===0);
     }'''.replace('CARD_HELPERS',CARD_HELPERS),arg=target,timeout=3000)
     page.wait_for_timeout(700)
     self.snapshot(page)
     left=[x for x in self.pending_controls(page) if x['id']==target['id'] and x['section']==target['section'] and x['kind']==target['kind']]
     diagnostic=dict(**target,attempt=attempt,control=detail,before=before,after=self.control_evidence(page,target),result='HALEN_KAPALI' if left else 'ACILDI')
     self.diagnostics.append(diagnostic)
     logging.info('NORMAL TIKLAMA | %s | %s | %s | deneme %s/3 | %s',target['name'],target['section'],target['kind'],attempt,diagnostic['result'])
     if not left:break
    except Exception as e:
     if isinstance(e,TaskError):raise
     # DOM tiklama sirasinda yenilendiyse once kalan kontrolu tekrar dogrula.
     self.snapshot(page)
     left=[x for x in self.pending_controls(page) if x['id']==target['id'] and x['section']==target['section'] and x['kind']==target['kind']]
     self.diagnostics.append(dict(**target,attempt=attempt,control=detail,before=before,after=self.control_evidence(page,target),result=safe_error(e),still_pending=bool(left)))
     if not left:break
     logging.warning('NORMAL TIKLAMA | %s | deneme %s/3 | %s',target['name'],attempt,safe_error(e))
    if attempt<3:page.wait_for_timeout(1000)
   else:
    self.snapshot(page)
    try:page.screenshot(path=str(self.checkpoint_dir/'metin_dugmesi.png'))
    except Exception as e:logging.warning('DUGME EKRANI KAYDEDILEMEDI | %s',safe_error(e))
    failed.append(target)
  self.checkpoint()
  if failed:
   detail=' | '.join(f'{t["name"]}, {t["section"]}, {t["kind"]}, ID={t["id"]}' for t in failed[:5])
   raise TaskError('KALAN_DUGME_ACILAMADI','METIN_GENISLET',f'{len(failed)} hedef uc denemede acilamadi; diger hedefler de denendi. '+detail)
 def report(self):
  rows=list(self.rows.values());total=self.meta.get('total');final=self.final_meta.get('total')
  checked=sum(r['quality']=='ALAN KONTROLU TAM' for r in rows)
  hist=dict(Counter(r['stars'] for r in rows if r['stars'] is not None));target={int(k):v for k,v in self.meta.get('histogram',{}).items()}
  coverage_ok=total is not None and total==final==len(rows)
  matched=coverage_ok and not self.error
  hist_ok=(len(target)==5 and all(hist.get(k,0)==v for k,v in target.items())) if target else None
  if coverage_ok and total==0:hist_ok=True
  missing=total-len(rows) if total is not None and len(rows)<=total and (final is None or total==final) else None
  return dict(total=total,final_total=final,unique=len(rows),fields_checked=checked,fields_incomplete=len(rows)-checked,
   text=sum(bool(r['text'] or r['translation']) for r in rows),stars_only=sum(r['comment_type'].startswith('YALNIZCA') for r in rows),
   missing=missing,distribution=hist,histogram_ok=hist_ok,
   coverage=('KAYIT SAYISI VE YILDIZ DAGILIMI ESLESTI' if hist_ok else 'KAYIT SAYISI ESLESTI; YILDIZ DAGILIMI OKUNAMADI' if hist_ok is None else 'YILDIZ DAGILIMI ESLESMEDI') if coverage_ok else 'KAYIT KAPSAMI DOGRULANAMADI',
   content='GORUNEN ALAN KONTROLLERI TAM; %100 GARANTI DEGIL' if checked==len(rows) and rows else 'METIN TAMLIGI DOGRULANAMAYAN KAYITLAR VAR',
   status='SAYI VE ALAN KONTROLLERI ESLESTI; %100 GARANTI DEGIL' if matched and hist_ok and checked==len(rows) else 'KISMI / TAMLIK DOGRULANAMADI')

def old_output_name(filename,name):
 return bool(re.fullmatch(re.escape(name)+r'_Yorumlar_\d{8}_\d{6}(?:_(?:KISMI|KONTROL_EDILDI))?\.(?:xlsx|html)',filename))

def cleanup_old_outputs(folder,name):
 # Yalniz bu uygulamanin ayni isletme icin urettigi eski dosyalar; alt klasor yok.
 for path in folder.iterdir():
  if path.is_file() and old_output_name(path.name,name):
   try:path.unlink();logging.info('ESKI CIKTI TEMIZLENDI | %s',path.name)
   except OSError as e:logging.warning('ESKI CIKTI SILINEMEDI | %s | %s',path.name,safe_error(e))

def verify_excel(path,expected):
 from openpyxl import load_workbook
 actual=load_workbook(path,read_only=True,data_only=False)
 try:
  if actual.sheetnames!=expected.sheetnames:raise TaskError('EXCEL_DOGRULAMA','EXCEL_KAYIT','Excel sayfalari eslesmiyor.')
  for left in expected:
   right=actual[left.title]
   if (left.max_row,left.max_column)!=(right.max_row,right.max_column):raise TaskError('EXCEL_DOGRULAMA','EXCEL_KAYIT',left.title+': satir/sutun sayisi eslesmiyor.')
   for lrow,rrow in zip(left.iter_rows(),right.iter_rows()):
    for a,b in zip(lrow,rrow):
     same=(a.value if a.value is not None else '')==(b.value if b.value is not None else '')
     # Excel ondalikli oranlari yaklasik 15-16 anlamli basamakla saklar.
     # Metinler ve tam sayilar birebir; yalniz float degerler kucuk toleransla.
     if not same and isinstance(a.value,float) and isinstance(b.value,(int,float)):
      same=math.isclose(a.value,b.value,rel_tol=1e-12,abs_tol=1e-12)
     if not same:
      raise TaskError('EXCEL_DOGRULAMA','EXCEL_KAYIT',left.title+'!'+a.coordinate+': hucre verisi eslesmiyor.')
     if isinstance(a.value,str) and a.value and b.data_type!='s':
      raise TaskError('EXCEL_DOGRULAMA','EXCEL_KAYIT',left.title+'!'+a.coordinate+': metin hucre turu eslesmiyor.')
 finally:actual.close()

MANAGER_HEADERS=['Yorum yapan kişi','Müşteri yorumu','Yıldız','İşletmenin cevabı','Yorum zamanı','Dil','Türkçe çeviri','Cevap zamanı']

def turkish_date(value):
 text=str(value or '').strip();prefix=''
 for word,label in [('edited ','Düzenlendi: '),('updated ','Güncellendi: ')]:
  if text.lower().startswith(word):prefix=label;text=text[len(word):];break
 special={'just now':'Az önce','today':'Bugün','yesterday':'Dün'}
 if text.lower() in special:return prefix+special[text.lower()]
 m=re.fullmatch(r'(\d+|a|an)\s+(minutes?|hours?|days?|weeks?|months?|years?)\s+ago',text,re.I)
 if m:
  amount=m[1] if m[1].isdigit() else '1'
  unit={'minute':'dakika','hour':'saat','day':'gün','week':'hafta','month':'ay','year':'yıl'}[m[2].lower().rstrip('s')]
  return prefix+f'{amount} {unit} önce'
 return prefix+text

def manager_values(row):
 checks=row.get('checks',{})
 comment=row.get('text') or row.get('translation') or ''
 if not comment:
  comment='Yorum yazılmamış; yalnızca yıldız verilmiş' if row.get('comment_type','').startswith('YALNIZCA') else 'Yorum metni alınamadı'
 elif not row.get('text'):comment+='\n[Gösterilen çeviri; özgün yorum alınamadı]'
 if checks.get('review_pending') and (row.get('text') or row.get('translation')):comment+='\n[Devamı alınamadı]'
 reply=row.get('reply') or row.get('reply_translation') or ''
 if not reply:
  reply='Yanıt yok' if row.get('reply_state')=='BU GORUNUMDE ISLETME YANITI YOK' else 'Yanıt bilgisi alınamadı'
 elif not row.get('reply'):reply+='\n[Gösterilen çeviri; özgün cevap alınamadı]'
 if checks.get('reply_pending') and (row.get('reply') or row.get('reply_translation')):reply+='\n[Devamı alınamadı]'
 code=row.get('language','').lower().split('-')[0]
 language={'tr':'Türkçe','en':'İngilizce','de':'Almanca','fr':'Fransızca','ar':'Arapça','ru':'Rusça','es':'İspanyolca','it':'İtalyanca'}.get(code,code or ('Metin yok' if row.get('comment_type','').startswith('YALNIZCA') else 'Belirlenemedi'))
 translation='Türkçe yazılmış' if code=='tr' else 'Metin yok' if row.get('comment_type','').startswith('YALNIZCA') else 'Çeviri yapılmadı'
 # Yeni model/API kurulmaz. Yalniz dili dogrulanmis mevcut Turkce ceviri gosterilebilir.
 if row.get('translation_language')=='tr' and row.get('translation'):translation=row['translation']
 reply_date=turkish_date(row.get('reply_date')) or ('Yanıt yok' if reply=='Yanıt yok' else 'Tarih alınamadı')
 return [row.get('name') or 'Ad alınamadı',comment,row['stars'] if row.get('stars') is not None else 'Alınamadı',reply,turkish_date(row.get('date')) or 'Tarih alınamadı',language,translation,reply_date]

def rounded_score(value):
 return float(Decimal(str(value)).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP))

def deep_analysis(rows):
 # Tarihler kaynakta yuvarlatilmis goreli ifadelerdir; takvim gunu uretmeyiz.
 periods=[('1 aydan kısa',0,30.436875),('1–3 ay önce',30.436875,91.310625),
  ('3–6 ay önce',91.310625,182.62125),('6–12 ay önce',182.62125,365.2425),
  ('1 yıl veya daha önce',365.2425,float('inf'))]
 checked=[r for r in rows if r.get('quality')=='ALAN KONTROLU TAM' and r.get('stars') in (1,2,3,4,5)]
 def stats(group):
  answered=sum(bool(r.get('reply') or r.get('reply_translation')) for r in group)
  missing=sum(r.get('reply_state')=='BU GORUNUMDE ISLETME YANITI YOK' for r in group)
  low=sum(r['stars']<=2 for r in group)
  return dict(count=len(group),average=sum(r['stars'] for r in group)/len(group) if group else None,
   low=low,low_rate=low/len(group) if group else None,answered=answered,unanswered=missing,
   unknown_reply=len(group)-answered-missing,answer_rate=answered/len(group) if group else None)
 groups=[]
 for label,start,end in periods:
  group=[r for r in checked if start<=age(r.get('date'))<end]
  groups.append(dict(period=label,**stats(group)))
 groups.append(dict(period='Tarih belirlenemedi',**stats([r for r in checked if age(r.get('date'))==float('inf')])))
 eligible=[g for g in groups[:-1] if g['count']>=5]
 comparison='Karşılaştırma için en az iki dönemde beşer değerlendirme gerekir.'
 if len(eligible)>=2:
  best=max(g['average'] for g in eligible);worst=min(g['average'] for g in eligible)
  if abs(best-worst)<1e-9:comparison='Karşılaştırılabilir dönemlerin ortalama puanları eşit.'
  else:
   labels=lambda value:', '.join(g['period'] for g in eligible if abs(g['average']-value)<1e-9)
   comparison=f"En yüksek ortalama: {labels(best)} ({rounded_score(best):.2f}/5). En düşük ortalama: {labels(worst)} ({rounded_score(worst):.2f}/5)."
 priorities=sorted([r for r in checked if r['stars']<=2 and r.get('reply_state')=='BU GORUNUMDE ISLETME YANITI YOK'],
  key=lambda r:(r['stars'],age(r.get('date'))))
 low=[r for r in checked if r['stars']<=2]
 return dict(total=len(rows),included=len(checked),excluded=len(rows)-len(checked),overall=stats(checked),
  low_rating=stats(low),periods=groups,comparison=comparison,minimum_period_count=5,
  priority_ids=[r['id'] for r in priorities],method='Yıldız, yanıt durumu ve yaklaşık göreli dönem; metin duygu analizi uygulanmadı.')

def add_analysis_sheet(wb,rows,analysis,started):
 from openpyxl.styles import Alignment,Font,PatternFill
 ws=wb.create_sheet('DERİN ANALİZ');ws.append(['Yönetici görünümü','Değer'])
 overview=analysis['overall'];low=analysis['low_rating']
 for label,value in [('Veri toplama zamanı',datetime.fromisoformat(started).strftime('%d.%m.%Y %H:%M')),
  ('Değerlendirme sayısı',analysis['total']),('Hesaplamaya alınan kayıt',analysis['included']),
  ('Çekilen kayıtların ortalama puanı',rounded_score(overview['average']) if overview['average'] is not None else 'Hesaplanamadı'),
  ('İşletme cevabı bulunan',overview['answered']),('Yanıt bulunmayan',overview['unanswered']),
  ('Yanıt durumu belirlenemeyen',overview['unknown_reply']),
  ('Yanıt oranı',overview['answer_rate'] if overview['answer_rate'] is not None else 'Hesaplanamadı'),
  ('1–2 yıldızlı değerlendirme',low['count']),('1–2 yıldızlı yanıtsız değerlendirme',low['unanswered'])]:ws.append([label,value])
 ws['B9'].number_format='0.0%'
 if analysis['excluded']:ws.append(['Hesaplamaya alınamayan kayıt',analysis['excluded']])
 ws.append(['Dönem karşılaştırması',analysis['comparison']])
 ws.append(['Okuma notu','Dönemler veri toplama zamanına göre yaklaşık yaş gruplarıdır; kesin takvim aralığı değildir. En yüksek/düşük ortalama yalnızca en az 5 kaydı olan dönemler içindir. Bu eşik istatistiksel anlamlılık göstermez; eski dönemlerin süreleri eşit değildir.'])
 ws.append(['Kapsam','Yıldız puanları müşteri değerlendirmesidir. Metin duygu analizi henüz uygulanmadı. Cevap bulunması sorunun çözüldüğünü göstermez.'])
 ws.append([])
 header=ws.max_row+1;ws.append(['Yaklaşık dönem','Kayıt sayısı','Ortalama yıldız','1–2 yıldız adedi','1–2 yıldız oranı','Cevap adedi','Cevap oranı','Karşılaştırma'])
 for g in analysis['periods']:
  ws.append([g['period'],g['count'],rounded_score(g['average']) if g['average'] is not None else 'Veri yok',g['low'],g['low_rate'] if g['low_rate'] is not None else 'Veri yok',g['answered'],g['answer_rate'] if g['answer_rate'] is not None else 'Veri yok',
   'Tarih bilinmiyor' if g['period']=='Tarih belirlenemedi' else 'Az veri' if g['count']<5 else 'Karşılaştırmaya dahil'])
  for col in (5,7):ws.cell(ws.max_row,col).number_format='0.0%'
  ws.cell(ws.max_row,3).number_format='0.00'
 ws.append([]);ws.append(['Öncelik: yanıtsız 1–2 yıldızlı değerlendirmeler'])
 priority_header=ws.max_row+1
 ws.append(['Yorum yapan kişi','Yıldız','Yorum zamanı','Yorumdan alıntı (ilk 500 karakter)','İşletmenin cevabı','Tam yoruma erişim'])
 ids=set(analysis['priority_ids']);lookup={r['id']:(i+2,r) for i,r in enumerate(rows)}
 for review_id in analysis['priority_ids']:
  source_row,r=lookup[review_id];values=manager_values(r);excerpt=values[1]
  if len(excerpt)>500:excerpt=excerpt[:500]+'… [Tam yorum için bağlantıyı açın]'
  ws.append([values[0],values[2],values[4],excerpt,'Yanıt yok','Tam yoruma git'])
  ws.cell(ws.max_row,6).hyperlink=f"#'TÜM YORUMLAR'!A{source_row}"
  ws.cell(ws.max_row,6).style='Hyperlink'
 if not ids:ws.append(['Yanıtsız 1–2 yıldızlı değerlendirme bulunmuyor.'])
 for index in (1,header,priority_header):
  for cell in ws[index]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245A81')
 for row in ws:
  for cell in row:
   cell.alignment=Alignment(wrap_text=True,vertical='top')
   if isinstance(cell.value,str):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    cell.value=ILLEGAL_CHARACTERS_RE.sub(lambda m:f'\\u{ord(m[0]):04x}',cell.value);cell.data_type='s'
  ws.row_dimensions[row[0].row].height=150 if row[0].row>priority_header else 90 if any(len(str(c.value or ''))>200 for c in row) else 36
 for row in ws.iter_rows(min_row=2,max_row=header-2):
  if isinstance(row[1].value,str) and len(row[1].value)>60:ws.merge_cells(start_row=row[0].row,start_column=2,end_row=row[0].row,end_column=8)
 for col,width in [('A',34),('B',12),('C',20),('D',75),('E',24),('F',23),('G',18),('H',25)]:ws.column_dimensions[col].width=width
 ws.freeze_panes='A2';ws.sheet_view.zoomScale=80
 ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0

def export(c,folder):
 from openpyxl import Workbook
 from openpyxl.styles import Alignment,Font,PatternFill
 from openpyxl.utils import get_column_letter
 from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
 folder.mkdir(parents=True,exist_ok=True)
 rows=sorted(c.rows.values(),key=lambda r:age(r['date']))
 values=[manager_values(r) for r in rows]
 for row in values:
  for i,v in enumerate(row):
   if isinstance(v,str):row[i]=ILLEGAL_CHARACTERS_RE.sub(lambda m:f'\\u{ord(m[0]):04x}',v)
 extras=[]
 for i,header in enumerate(MANAGER_HEADERS):
  maximum=max([len(r[i]) for r in values if isinstance(r[i],str)] or [0])
  extras.extend((i,p,header+f' (devam {p})') for p in range(1,(maximum-1)//32767+1))
 wb=Workbook();summary=wb.active;summary.title='ÖZET';summary.append(['Bilgi','Değer'])
 report=c.report()
 for label,value in [('İşletme adı',c.meta.get('name','')),('Google Maps bağlantısı',c.url),('Ortalama puan',c.meta.get('rating')),('Değerlendirme sayısı',report['total'])]:summary.append([label,value if value is not None else 'Alınamadı'])
 histogram={int(k):v for k,v in c.meta.get('histogram',{}).items()}
 for star in range(1,6):summary.append([f'{star} yıldız',histogram.get(star,'Alınamadı')])
 if report['total']!=len(rows):summary.append(['Exceldeki değerlendirme sayısı',len(rows)])
 summary.append(['Rapor zamanı',datetime.fromisoformat(c.started).strftime('%d.%m.%Y %H:%M')])
 summary['B4'].number_format='0.0'
 ws=wb.create_sheet('TÜM YORUMLAR');ws.append(MANAGER_HEADERS+[x[2] for x in extras])
 for row in values:ws.append([v[:32767] if isinstance(v,str) else v for v in row]+[row[i][p*32767:(p+1)*32767] for i,p,_ in extras])
 for sheet in wb:
  sheet.freeze_panes='A2'
  for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245A81')
  for row in sheet.iter_rows(min_row=2):
   for cell in row:
    cell.alignment=Alignment(wrap_text=True,vertical='top')
    if isinstance(cell.value,str):cell.data_type='s'
 if c.url.startswith('https://'):summary['B3'].hyperlink=c.url;summary['B3'].style='Hyperlink'
 summary.column_dimensions['A'].width=28;summary.column_dimensions['B'].width=75
 ws.auto_filter.ref=ws.dimensions
 for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width={1:24,2:52,3:8,4:52,5:20,6:16,7:45,8:20}.get(col,52)
 for row in range(2,ws.max_row+1):
  lines=max(sum(max(1,(len(part)+51)//52) for part in str(ws.cell(row,col).value or '').split('\n')) for col in (2,4))
  ws.row_dimensions[row].height=min(409,max(48,lines*15))
 if getattr(c,'include_analysis',False):
  c.analysis=deep_analysis(rows)
  add_analysis_sheet(wb,rows,c.analysis,c.started)
 name=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',c.meta.get('name') or 'Google_Maps').strip(' .')[:100]
 path=folder/(name+'_Yorumlar.xlsx')
 fd,tmp=tempfile.mkstemp(suffix='.xlsx',dir=folder);os.close(fd)
 try:
  wb.save(tmp);c.progress.update('Excel dosyası doğrulanıyor',97);verify_excel(tmp,wb)
  try:os.replace(tmp,path)
  except PermissionError as e:raise TaskError('EXCEL_ACIK','EXCEL_KAYIT','Excel dosyasi acik/kilitli. Excel penceresini kapatip yeniden deneyin; onceki dosya korundu.') from e
 finally:
  if os.path.exists(tmp):os.unlink(tmp)
 cleanup_old_outputs(folder,name)
 return path

def save_run_result(c,files=None):
 result=dict(version=VERSION,started=c.started,report=c.report(),error=c.error,diagnostics=c.diagnostics,access_state=c.access_state,access_screenshot=c.access_screenshot,progress=c.progress.history,performance=c.performance,
  manager_layout=dict(headers=MANAGER_HEADERS,row_ids=[r['id'] for r in sorted(c.rows.values(),key=lambda r:age(r['date']))]),
  files=dict(checkpoint=str(c.checkpoint_file),log=str(c.checkpoint_dir/'calisma.log'),**(files or {})))
 path=c.checkpoint_dir/'sonuc.json';tmp=path.with_suffix('.tmp')
 if hasattr(c,'analysis'):result['analysis']=dict(source=getattr(c,'analysis_source',None),**c.analysis)
 tmp.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');tmp.replace(path)
 return path

def summary(c):
 r=c.report();print('\nSonuç')
 print('Toplanan değerlendirme:',r['unique'])
 print('Görünen toplam:',r['total'] if r['total'] is not None else 'Alınamadı')
 print('Metinli yorum:',r['text'],'| Yalnızca yıldız:',r['stars_only'])
 print('İşletme cevabı:',sum(bool(x['reply'] or x['reply_translation']) for x in c.rows.values()))
 print('Sonuç:', 'İncelenmesi gereken noktalar var.' if r['status'].startswith('KISMI') else 'Kayıt ve metin kontrolleri eşleşti.')

def self_test():
 assert metadata({'counts':['105 reviews','4.3105 reviews','543214.3105 reviews']})['total']==105
 assert metadata({'counts':['1,234 reviews']})['total']==1234
 raw=dict(id='A',name='Ayni',star_label='1 star',date='a month ago',text='',reply='Yanit',has_reply=True)
 assert normalize(raw)['comment_type'].startswith('YALNIZCA')
 raw.update(text='Ceviri',translated=True,reply_translated=True)
 row=normalize(raw);assert not row['text'] and row['translation']=='Ceviri' and not row['reply'] and row['reply_translation']=='Yanit'
 assert age('10 hours ago')<age('6 days ago')<age('11 months ago')<age('a year ago')
 assert age('3 hafta önce')==21
 # Yazar noktalamasini veri kaybi diye etiketleme; kapali devam dugmesini yine yakala.
 raw=dict(id='A',name='Ayni',star_label='5 stars',date='2 hours ago',has_text=True,has_reply=True,
  text='Yorum...\nIkinci satir…',reply='Yanit...',reply_date='today')
 row=normalize(raw);assert row['quality']=='ALAN KONTROLU TAM' and row['text']==raw['text'] and row['reply']==raw['reply']
 raw['review_pending']=1
 assert normalize(raw)['issues']=='yorum genisletilemedi'
 raw['review_pending']=0;raw['translation_unknown']=True
 assert 'ceviri dugmesinin' in normalize(raw)['issues']
 c=Collector.__new__(Collector);c.rows={'A':row};c.meta={'total':1,'histogram':{'1':0,'2':0,'3':0,'4':0,'5':1}};c.final_meta={'total':1};c.error=None
 assert c.report()['histogram_ok'] is True and not c.report()['status'].startswith('KISMI')
 c.meta['histogram']={}
 assert c.report()['histogram_ok'] is None and 'OKUNAMADI' in c.report()['coverage']
 assert is_limited_notice("You're seeing a limited view of Google Maps")
 assert is_limited_notice('Google Haritalar’ın sınırlı bir görünümünü kullanıyorsunuz')
 assert not is_limited_notice('Learn more about using limited view')
 assert not is_limited_notice('Sınırlı görünüm hakkında bilgi edinin')
 assert old_output_name('FAMİ_Yorumlar_20261013_171839_KISMI.xlsx','FAMİ')
 assert not old_output_name('FAMİ_Yorumlar.xlsx','FAMİ')
 assert not old_output_name('Başka_Yorumlar_20261013_171839_KISMI.xlsx','FAMİ')
 assert unique_control({'card_count':2,'control_count':1})
 assert not unique_control({'card_count':2,'control_count':2})
 assert not unique_control({'card_count':1,'control_count':0})
 assert not unique_control({'card_count':0,'control_count':0})
 # Ayni yorumu temsil eden yanitsiz kopya, diger kopyanin yanitini silemez.
 primary=dict(id='A',name='Kisi',star_label='5 stars',date='today',has_text=True,text='Yorum...',has_reply=False,review_pending=0)
 with_reply=dict(primary,has_reply=True,reply='Isletmenin tam yaniti...',reply_date='today',reply_pending=0,all_text='Kisi\nYorum...\nIsletmenin tam yaniti...')
 merged=merge_raw_variants([with_reply,primary]);assert merged['reply']==with_reply['reply'] and merged['reply_date']=='today'
 assert normalize(merged)['reply_state']=='YANIT ACIK / OKUNDU'
 assert normalize(primary)['reply_state']=='BU GORUNUMDE ISLETME YANITI YOK'
 c=Collector.__new__(Collector);c.rows={};c.raw_rows={};c.store(with_reply);c.store(primary)
 assert c.rows['A']['reply']==with_reply['reply'] and c.rows['A']['quality']=='EKSIK / DOGRULANAMADI'
 conflict=merge_raw_variants([primary,dict(primary,text='Farkli bir yorum')])
 assert conflict['variant_conflicts']==['text'] and normalize(conflict)['quality']=='EKSIK / DOGRULANAMADI'
 shown=manager_values(normalize(primary));assert len(shown)==8 and shown[3]=='Yanıt yok' and shown[7]=='Yanıt yok'
 assert 'yalnızca yıldız' in manager_values(normalize(dict(primary,text='',has_text=False)))[1]
 unknown=normalize(dict(primary,has_reply=True,reply='',reply_pending=1))
 assert manager_values(unknown)[3]=='Yanıt bilgisi alınamadı'
 english=normalize(dict(primary,language='en'))
 assert manager_values(english)[5:7]==['İngilizce','Çeviri yapılmadı']
 assert manager_values(normalize(dict(primary,review_pending=1)))[1].endswith('[Devamı alınamadı]')
 for s in ['https://example.com/maps','https://www.google.com/search?q=test']:
  try:validate_url(s);raise AssertionError('Invalid URL accepted')
  except TaskError:pass
 assert turkish_date('2 months ago')=='2 ay önce'
 assert turkish_date('Edited a year ago')=='Düzenlendi: 1 yıl önce'
 assert remaining_text(125)=='2 dakika 5 saniye' and remaining_text(None)=='hesaplanıyor'
 assert rounded_score(3.625)==3.63 and age('yesterday')==1 and age('dün')==1
 sample=lambda i,date,stars=5,reply='':dict(id=str(i),date=date,stars=stars,reply=reply,quality='ALAN KONTROLU TAM',reply_state='YANIT ACIK / OKUNDU' if reply else 'BU GORUNUMDE ISLETME YANITI YOK')
 boundary=deep_analysis([sample(i,d) for i,d in enumerate(['today','1 month ago','3 months ago','6 months ago','a year ago','unknown'])])
 assert [g['count'] for g in boundary['periods']]==[1,1,1,1,1,1]
 assert 'en az iki' in boundary['comparison'] and boundary['overall']['answer_rate']==0
 mixed=deep_analysis([sample(1,'a month ago',1),sample(2,'today',2,'Cevap'),dict(sample(3,'today',1),quality='EKSIK / DOGRULANAMADI')])
 assert mixed['included']==2 and mixed['excluded']==1 and mixed['priority_ids']==['1'] and mixed['low_rating']['answered']==1
 assert deep_analysis([])['overall']['average'] is None
 # Emmi: gecerli gruplama gibi gorunen birlesik puan+toplam histograma uymali.
 hist_labels=['5 stars, 9,777 reviews','4 stars, 3,348 reviews','3 stars, 1,715 reviews','2 stars, 767 reviews','1 stars, 1,674 reviews']
 meta=metadata(dict(counts=['17,281 reviews','4.117,281 reviews','543214.117,281 reviews'],labels=hist_labels))
 assert meta['total']==17281 and meta['histogram_total']==17281 and 'dagilimi' in meta['total_source']
 assert metadata(dict(counts=[],labels=hist_labels))['total_source']=='yildiz dagilimi'
 assert metadata(dict(counts=['107 reviews','108 reviews'],labels=[]))['total'] is None
 assert metadata(dict(counts=['17.281 yorum'],labels=[]))['total']==17281
 assert metadata(dict(counts=['17\u202f281 reviews'],labels=[]))['total']==17281
 assert metadata(dict(counts=['17K reviews'],labels=[]))['total'] is None
 assert map_place_key('https://www.google.com/maps/place/X/data=!4m2!3m1!1s0x123:0x456')=='0x123:0x456'
 assert map_place_key('https://www.google.com/maps/place/Honda+Plaza+Damla/@1,2,3z')=='name:honda plaza damla'
 from unittest.mock import patch
 from contextlib import redirect_stdout
 import io
 buffer=io.StringIO();progress=Progress()
 with patch('time.monotonic',side_effect=[0,10,20,30]),redirect_stdout(buffer):
  progress.update('Değerlendirmeler yükleniyor',10,60,done=0,total=100)
  progress.update('Değerlendirmeler yükleniyor',10,60,done=20,total=100)
  progress.update('Değerlendirmeler yükleniyor',10,60,done=20,total=100,note='Yeni kayıt bekleniyor')
  progress.update('Değerlendirmeler yükleniyor',10,60,done=30,total=100)
 assert progress.percent==25 and progress.history[-1]['remaining_seconds']==70
 with redirect_stdout(buffer):Progress().update('Değerlendirmeler yükleniyor',10,60,done=1020,total=None)
 assert '1020 kayıt alındı; toplam okunamadı' in buffer.getvalue()
 print('Yerel kontroller başarılı.')

def analyze_saved(source=None):
 root=WORKSPACE/'Calisma_kayitlari'
 if source:path=Path(source).expanduser().resolve()
 else:
  candidates=sorted((p/'yorum_kurtarma.json' for p in root.iterdir() if p.is_dir() and re.fullmatch(r'\d{8}_\d{6}_\d+',p.name) and (p/'yorum_kurtarma.json').is_file()),reverse=True) if root.exists() else []
  if not candidates:print('Kayıtlı yorum verisi bulunamadı. Önce yorum çekme işlemini çalıştırın.');return 2
  path=candidates[0]
 try:
  data=json.loads(path.read_text(encoding='utf-8'))
  rows=data['rows']
  if not rows or len({r['id'] for r in rows})!=len(rows):raise ValueError('Kayıtlar boş veya kimlikler tekrarlı.')
  directory=root/('Analiz_'+datetime.now().strftime('%Y%m%d_%H%M%S_%f'))
  c=Collector(validate_url(data['url']),run_dir=directory)
  for field in ('meta','final_meta','error','access_state','access_screenshot','diagnostics','started','round'):
   if field in data:setattr(c,field,data[field])
  c.rows={r['id']:r for r in rows};c.checkpoint_file=path
  c.include_analysis=True
  c.analysis_source=dict(path=str(path),version=data.get('version'),started=data['started'])
  if c.report()['status'].startswith('KISMI'):
   print('Son kaydın tamlığı doğrulanamamış. Mevcut Excel değiştirilmedi; önce kaynak çalışma incelenmeli.');return 2
  # Kaynak kurtarma JSON'u degistirilmez; analiz kaydi ayri tutulur.
  logging.basicConfig(level=logging.INFO,handlers=[logging.FileHandler(directory/'calisma.log',encoding='utf-8')],force=True)
  if importlib.util.find_spec('openpyxl') is None:
   print('openpyxl gerekli: python -m pip install openpyxl');return 3
  print('Kayıtlı yorumlar analiz ediliyor. Kaynak zamanı:',datetime.fromisoformat(c.started).strftime('%d.%m.%Y %H:%M'))
  c.progress.update('Derin analiz ve Excel hazırlanıyor',95)
  excel=export(c,DEFAULT_OUTPUT);save_run_result(c,dict(excel=str(excel)))
  c.progress.update('Analiz tamamlandı',100,force=True)
  save_run_result(c,dict(excel=str(excel)))
  print('Excel güncellendi:',excel);print('Analiz kaydı:',directory/'sonuc.json');return 0
 except (Exception,KeyboardInterrupt) as e:
  logging.error('ANALIZ BASARISIZ | %s',safe_error(e))
  print('Analiz kaydedilemedi:',safe_error(e));return 3

def main():
 for s in (sys.stdout,sys.stderr):
  if hasattr(s,'reconfigure'):s.reconfigure(encoding='utf-8',errors='replace')
 if '--self-test' in sys.argv:self_test();return 0
 if '--analyze' in sys.argv:
  if len(sys.argv) not in (2,3) or sys.argv[1]!='--analyze':print('Kullanım: python google_yorumlari_otomatik.py --analyze [yorum_kurtarma.json]');return 3
  return analyze_saved(sys.argv[2] if len(sys.argv)==3 else None)
 run_dir=WORKSPACE/'Calisma_kayitlari'/datetime.now().strftime('%Y%m%d_%H%M%S_%f')
 run_dir.mkdir(parents=True,exist_ok=True)
 logging.basicConfig(level=logging.INFO,format='%(asctime)s | %(levelname)s | %(message)s',
  handlers=[logging.FileHandler(run_dir/'calisma.log',encoding='utf-8')],force=True)
 print('Google Maps yorumları | Sürüm',VERSION)
 print('Genel yüzde yaklaşık aşama ilerlemesidir; süre tahmini ölçülen hıza göre güncellenir.')
 logging.info('BASLANGIC | SURUM=%s | PYTHON=%s | WORKSPACE=%s',VERSION,sys.version.split()[0],WORKSPACE)
 if sys.version_info<(3,10):print('Python 3.10 veya üzeri gerekli.');return 3
 try:url=validate_url(input('Google Maps işletme bağlantısını yapıştırın: ').strip().strip('"'))
 except (TaskError,EOFError) as e:logging.error('LINK HATASI | %s',safe_error(e));print('Geçerli bir Google Maps işletme bağlantısı girilmedi.');return 3
 c=Collector(url,run_dir=run_dir)
 c.progress.update('Hazırlık yapılıyor',0)
 try:
  bootstrap()
  from playwright.sync_api import sync_playwright
  with sync_playwright() as pw:
   c.progress.update('Chrome hazırlanıyor',5)
   browser=chrome(pw)
   context=browser.contexts[0];context.set_default_timeout(15000)
   page=context.new_page()
   expected_key=map_place_key(url)
   retry('MAPS_AC',lambda:page.goto(url,wait_until='domcontentloaded',timeout=60000))
   resolved_key=map_place_key(page.url)
   if resolved_key and (not expected_key or expected_key.startswith('name:')):expected_key=resolved_key
   print('\nChrome açıldı. Google girişini ve çerez tercihini tarayıcıda tamamlayın.')
   print('Şifrenizi konsola yazmayın. İşletme sayfasına dönün.')
   input('İşletme sayfası hazır olduğunda ENTER tuşuna basın: ')
   c.progress.update('İşletme bilgileri okunuyor',10)
   # Kullanicinin acmis olabilecegi diger Maps sekmesini kullan.
   pages=[p for p in context.pages if not p.is_closed() and '/maps' in p.url and 'google.' in (urlsplit(p.url).hostname or '')]
   if expected_key:
    matches=[p for p in pages if map_place_key(p.url)==expected_key]
    if not matches:raise TaskError('HEDEF_ISLETME_ACIK_DEGIL','GIRIS','Girilen baglantidaki isletme sekmesi acik degil; baska isletmeye gecilmis olabilir.')
    if page not in matches:page=matches[-1]
   elif page.is_closed() or '/maps' not in page.url:
    raise TaskError('MAPS_ACIK_DEGIL','GIRIS','Programin actigi Google Maps isletme sekmesine donun; eski sekme otomatik secilmedi.')
   page.bring_to_front()
   collection_ok=False
   try:c.collect(page);collection_ok=True
   finally:
    # Genisletme hatasi olsa da son toplam ve son DOM durumu okunabiliyorsa korunur.
    try:
     c.check_access(page,allow_limited=True)
     c.final_meta=retry('SON_TOPLAM_KONTROLU',lambda:metadata(page.evaluate(READ_META)))
     c.snapshot(page)
     if collection_ok:c.progress.update('Son kayıtlar doğrulanıyor',92)
     logging.info('SON KONTROL | MAPS=%s | BENZERSIZ=%s',c.final_meta.get('total'),len(c.rows))
    except Exception as final_error:
     logging.warning('SON KONTROL YAPILAMADI | %s',safe_error(final_error))
 except (Exception,KeyboardInterrupt) as e:
  tb=traceback.extract_tb(e.__traceback__);loc=f'{Path(tb[-1].filename).name}:{tb[-1].lineno}' if tb else ''
  c.error=dict(code=getattr(e,'code','KULLANICI_DURDURDU' if isinstance(e,KeyboardInterrupt) else 'BEKLENMEYEN_HATA'),stage=getattr(e,'stage','ANA_ISLEM'),reason=safe_error(e),evidence=getattr(e,'evidence',None),location=loc,round=c.round,records=len(c.rows))
  logging.error('ISLEM DURDU | %s',json.dumps(c.error,ensure_ascii=False));c.checkpoint()
  explanations={
   'ILK_YORUM_YUKLENMEDI':'Yorum paneli açıldı ancak ilk yorumlar yüklenmedi.',
   'YORUM_GIRISI_BULUNAMADI':'İşletmenin yorum sekmesi veya değerlendirme düğmesi bulunamadı.',
   'YORUM_SEKMESI_ACILMADI':'Yorum sekmesi seçilemedi.',
   'YORUM_PANELI_BELIRSIZ':'Doğru işletme paneli belirlenemedi.',
   'HEDEF_ISLETME_ACIK_DEGIL':'Girilen bağlantının işletme sekmesi açık değil. Hedef işletmeye dönüp yeniden çalıştırın.',
   'YORUM_YUKLEME_ERKEN_DURDU':'Yorumların yüklenmesi görünen toplamdan önce durdu. Alınan kayıtlar korunuyor.',
   'KULLANICI_DURDURDU':'İşlem kullanıcı tarafından durduruldu.',
   'SURE_SINIRI':'İki saatlik çalışma sınırına ulaşıldı. Alınan kayıtlar korunuyor.'}
  print(explanations.get(c.error['code'],'Toplama tamamlanamadı. Ayrıntılar çalışma kaydına yazıldı.'))
  print('Alınan kayıt:',len(c.rows),'| Görünen toplam:',c.meta.get('total') if c.meta.get('total') is not None else c.final_meta.get('total','Okunamadı'))
 save_run_result(c)
 if not c.rows and (c.meta.get('total')!=0 or c.error):print('Değerlendirme alınamadı. Ayrıntılar çalışma kaydında:',run_dir);return 2
 if not c.rows:print('İşletmede değerlendirme görünmüyor; işletme özeti kaydediliyor.')
 try:
  folder=DEFAULT_OUTPUT
  c.progress.update('Excel hazırlanıyor',95)
  path=retry('EXCEL_KAYDET',lambda:export(c,folder))
  c.progress.update('Excel kaydedildi; sonuç kısmi' if c.report()['status'].startswith('KISMI') else 'İşlem tamamlandı',100,force=True)
  print('Excel kaydedildi:',path)
  summary(c)
  c.checkpoint()
  save_run_result(c,dict(excel=str(path)))
  logging.info('KAYIT BITTI | EXCEL=%s | DURUM=%s',path,c.report()['status'])
  if c.report()['status'].startswith('KISMI'):print('Ayrıntılı inceleme kaydı:',run_dir)
 except (Exception,KeyboardInterrupt) as e:
  logging.error('KAYIT BASARISIZ | %s',safe_error(e));print('Excel kaydedilemedi. Açık Excel penceresini ve kayıt klasörünü kontrol edin.');print('Kurtarma kaydı:',c.checkpoint_file);return 3
 return 2 if c.report()['status'].startswith('KISMI') else 0

if __name__=='__main__':
 try:sys.exit(main())
 except Exception as e:
  logging.error('BAŞLATMA HATASI | %s',safe_error(e));print('Program başlatılamadı. Ayrıntılı çalışma kaydını kontrol edin.');sys.exit(3)
